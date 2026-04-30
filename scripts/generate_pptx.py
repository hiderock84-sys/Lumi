#!/usr/bin/env python3
"""JSON定義からPowerPoint(.pptx)を自動生成するスクリプト。"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.util import Pt


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Create a PowerPoint file from a JSON outline."
    )
    parser.add_argument(
        "--input",
        required=True,
        help="Path to input JSON file.",
    )
    parser.add_argument(
        "--output",
        help="Output .pptx path. If omitted, JSONのoutput値またはinput名を利用します。",
    )
    return parser.parse_args()


def load_json(input_path: Path) -> dict[str, Any]:
    with input_path.open("r", encoding="utf-8") as f:
        data = json.load(f)

    if "slides" not in data or not isinstance(data["slides"], list):
        raise ValueError("Input JSON must contain a list field named 'slides'.")
    return data


def resolve_output_path(args_output: str | None, data: dict[str, Any], input_path: Path) -> Path:
    if args_output:
        return Path(args_output).expanduser().resolve()
    if isinstance(data.get("output"), str) and data["output"].strip():
        return Path(data["output"]).expanduser().resolve()
    return input_path.with_suffix(".pptx").resolve()


def _set_font_size(shape: Any, font_size_pt: int | None) -> None:
    if not font_size_pt or not getattr(shape, "has_text_frame", False):
        return
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(font_size_pt)


def _add_notes(slide: Any, notes: str | None) -> None:
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def _find_body_placeholder(slide: Any) -> Any | None:
    for placeholder in slide.placeholders:
        placeholder_type = placeholder.placeholder_format.type
        if placeholder_type in (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT):
            return placeholder
    return None


def add_title_slide(
    prs: Presentation,
    title: str,
    subtitle: str | None,
    notes: str | None,
    title_font_size: int | None,
    body_font_size: int | None,
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    if slide.shapes.title:
        slide.shapes.title.text = title
        _set_font_size(slide.shapes.title, title_font_size)

    if subtitle and len(slide.placeholders) > 1:
        subtitle_placeholder = slide.placeholders[1]
        subtitle_placeholder.text = subtitle
        _set_font_size(subtitle_placeholder, body_font_size)

    _add_notes(slide, notes)


def add_section_slide(
    prs: Presentation,
    title: str,
    subtitle: str | None,
    notes: str | None,
    title_font_size: int | None,
    body_font_size: int | None,
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[2])
    if slide.shapes.title:
        slide.shapes.title.text = title
        _set_font_size(slide.shapes.title, title_font_size)

    if subtitle and len(slide.placeholders) > 1:
        subtitle_placeholder = slide.placeholders[1]
        subtitle_placeholder.text = subtitle
        _set_font_size(subtitle_placeholder, body_font_size)

    _add_notes(slide, notes)


def add_bullet_slide(
    prs: Presentation,
    title: str,
    bullets: list[Any] | None,
    notes: str | None,
    title_font_size: int | None,
    body_font_size: int | None,
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    if slide.shapes.title:
        slide.shapes.title.text = title
        _set_font_size(slide.shapes.title, title_font_size)

    body_placeholder = _find_body_placeholder(slide)
    if body_placeholder and isinstance(bullets, list):
        text_frame = body_placeholder.text_frame
        text_frame.clear()

        for index, bullet in enumerate(bullets):
            paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
            if isinstance(bullet, dict):
                paragraph.text = str(bullet.get("text", ""))
                paragraph.level = int(bullet.get("level", 0))
            else:
                paragraph.text = str(bullet)
                paragraph.level = 0

        _set_font_size(body_placeholder, body_font_size)

    _add_notes(slide, notes)


def build_presentation(data: dict[str, Any]) -> Presentation:
    prs = Presentation()
    theme = data.get("theme", {})
    title_font_size = theme.get("title_font_size_pt")
    body_font_size = theme.get("body_font_size_pt")

    for slide_data in data["slides"]:
        slide_type = slide_data.get("type", "content")
        title = str(slide_data.get("title", "")).strip()
        subtitle = slide_data.get("subtitle")
        notes = slide_data.get("notes")
        bullets = slide_data.get("bullets")

        if not title:
            raise ValueError("Each slide must have a non-empty 'title'.")

        if slide_type == "title":
            add_title_slide(prs, title, subtitle, notes, title_font_size, body_font_size)
        elif slide_type == "section":
            add_section_slide(prs, title, subtitle, notes, title_font_size, body_font_size)
        elif slide_type == "content":
            add_bullet_slide(prs, title, bullets, notes, title_font_size, body_font_size)
        else:
            raise ValueError(
                "Unsupported slide type: "
                f"{slide_type}. Allowed: title, section, content."
            )

    return prs


def main() -> None:
    args = parse_args()
    input_path = Path(args.input).expanduser().resolve()
    data = load_json(input_path)
    output_path = resolve_output_path(args.output, data, input_path)

    presentation = build_presentation(data)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(output_path))
    print(f"Generated: {output_path}")


if __name__ == "__main__":
    main()
