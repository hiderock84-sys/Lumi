#!/usr/bin/env python3
"""参考スクリーンショット（第2回分）を全面リデザインしたPPTを生成する。

対象:
- 5ステージ制 回復プログラム
- 相模原ダルクグループ（8拠点）
- 回復プログラム（デイ/ナイト/個別）
- 依存症の理解（7つの特徴）
- 交差依存（クロスアディクション）
- 長期離脱症状（PAWS）
- 家族支援プログラム（CRAFT/家族会）
- 連携機関とネットワーク
- お問い合わせ

加えて、既存のPart1と結合した完全版（18枚）を同時生成する。
"""

from __future__ import annotations

import argparse
import copy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build redesigned part2 PPT and merged complete deck."
    )
    parser.add_argument(
        "--part1-input",
        default="/workspace/outputs/sagamihara_darc_redesign_part1_from_screenshots.pptx",
        help="Part1 pptx path for merged complete deck.",
    )
    parser.add_argument(
        "--output-part2",
        default="/workspace/outputs/sagamihara_darc_redesign_part2_from_screenshots.pptx",
        help="Output part2 pptx path.",
    )
    parser.add_argument(
        "--output-complete",
        default="/workspace/outputs/sagamihara_darc_redesign_complete_from_screenshots.pptx",
        help="Output merged complete pptx path.",
    )
    return parser.parse_args()


class Theme:
    bg = RGBColor(242, 246, 252)
    panel = RGBColor(226, 235, 248)
    card = RGBColor(255, 255, 255)
    navy = RGBColor(15, 33, 66)
    header = RGBColor(22, 47, 91)
    accent = RGBColor(0, 184, 179)
    accent2 = RGBColor(45, 112, 255)
    accent3 = RGBColor(255, 153, 61)
    text = RGBColor(22, 30, 47)
    sub = RGBColor(82, 96, 124)
    white = RGBColor(255, 255, 255)


def add_rect(slide, left, top, width, height, color: RGBColor, line: RGBColor | None = None):
    shp = slide.shapes.add_shape(1, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1.2)
    return shp


def add_text(
    slide,
    left,
    top,
    width,
    height,
    text: str,
    size: int = 16,
    bold: bool = False,
    color: RGBColor = Theme.text,
    align: PP_ALIGN = PP_ALIGN.LEFT,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    run = p.runs[0]
    run.font.name = "Meiryo"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_bullets(
    slide,
    left,
    top,
    width,
    height,
    items: list[str],
    size: int = 12,
    color: RGBColor = Theme.text,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"• {item}"
        p.level = 0
        if p.runs:
            run = p.runs[0]
            run.font.name = "Meiryo"
            run.font.size = Pt(size)
            run.font.color.rgb = color
    return box


def header(slide, title: str, subtitle: str, page: str):
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, Theme.bg)
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.2), Theme.accent)
    add_rect(slide, Inches(0), Inches(0.2), SLIDE_W, Inches(0.92), Theme.header)
    add_text(
        slide,
        Inches(0.45),
        Inches(0.39),
        Inches(8.9),
        Inches(0.35),
        title,
        size=20,
        bold=True,
        color=Theme.white,
    )
    add_text(
        slide,
        Inches(0.45),
        Inches(0.74),
        Inches(10.9),
        Inches(0.25),
        subtitle,
        size=11,
        color=RGBColor(198, 215, 242),
    )
    add_text(
        slide,
        Inches(12.1),
        Inches(0.42),
        Inches(1.0),
        Inches(0.25),
        page,
        size=11,
        color=Theme.white,
        align=PP_ALIGN.RIGHT,
    )
    add_rect(slide, Inches(0), Inches(7.2), SLIDE_W, Inches(0.3), Theme.panel)
    add_text(
        slide,
        Inches(0.35),
        Inches(7.26),
        Inches(8.5),
        Inches(0.2),
        "SAGAMIHARA DARC | STRATEGIC REDESIGN (PART 2)",
        size=9,
        color=Theme.sub,
    )


def cover(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, Theme.navy)
    add_rect(slide, Inches(0), Inches(0), Inches(5.9), SLIDE_H, RGBColor(11, 24, 49))
    add_rect(slide, Inches(5.9), Inches(0), Inches(7.433), SLIDE_H, RGBColor(28, 64, 121))
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.2), Theme.accent)
    add_text(
        slide,
        Inches(0.62),
        Inches(1.0),
        Inches(4.9),
        Inches(2.8),
        "相模原ダルク\n企業提案品質\nリデザイン資料",
        size=35,
        bold=True,
        color=Theme.white,
    )
    add_text(
        slide,
        Inches(0.62),
        Inches(4.35),
        Inches(5.0),
        Inches(1.6),
        "参考スクリーンショット（第2回分）の内容を\n省略せず再構築",
        size=14,
        color=RGBColor(180, 204, 240),
    )
    add_rect(slide, Inches(6.25), Inches(1.0), Inches(6.5), Inches(5.7), RGBColor(248, 251, 255))
    add_text(
        slide,
        Inches(6.55),
        Inches(1.28),
        Inches(5.9),
        Inches(0.5),
        "収録セクション（PAGE 10-18相当）",
        size=20,
        bold=True,
        color=Theme.navy,
    )
    add_bullets(
        slide,
        Inches(6.55),
        Inches(1.95),
        Inches(5.9),
        Inches(4.9),
        [
            "5ステージ制 回復プログラム",
            "相模原ダルクグループ（通所2 / 入寮6）",
            "回復プログラム（デイケア / ナイトケア / 個別）",
            "依存症の理解（7つの特徴）",
            "交差依存（クロスアディクション）",
            "長期離脱症状（PAWS）",
            "家族支援プログラム（CRAFT / 家族会）",
            "連携機関とネットワーク（医療・行政・司法・地域）",
            "お問い合わせ（CHANGE YOUR LIFE!）",
        ],
        size=13,
    )


def stage_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "5ステージ制 回復プログラム", "MEMBERからMANAGERまで段階的に自立を実装", "10")

    stages = [
        ("MEMBER", "導入期 1-3ヶ月", ["離脱症状安定", "24h見守り・服薬管理", "NA/AA・スポーツ"]),
        ("SUPPORT", "基礎期 4-6ヶ月", ["認知変容と役割理解", "SMARPP全24回（CBT）", "施設内係活動"]),
        ("TRAINEE", "成長期 7-12ヶ月", ["リーダーシップ発揮", "新入寮者サポート", "WRAP実践"]),
        ("CHIEF", "準備期 13-18ヶ月", ["就労先確保・通勤準備", "単独外出の拡大", "B型訓練"]),
        ("MANAGER", "自立期 18ヶ月〜", ["外部就労", "一人暮らし開始", "金銭・服薬自己管理"]),
    ]

    x = Inches(0.52)
    w = Inches(2.5)
    gap = Inches(0.1)
    for i, (name, period, bullets) in enumerate(stages):
        card_x = x + i * (w + gap)
        color = [Theme.accent2, Theme.accent, RGBColor(50, 180, 100), RGBColor(138, 101, 255), RGBColor(228, 89, 124)][i]
        add_rect(slide, card_x, Inches(1.45), w, Inches(5.45), Theme.card, Theme.panel)
        add_rect(slide, card_x, Inches(1.45), w, Inches(0.42), color)
        add_text(slide, card_x + Inches(0.12), Inches(1.53), Inches(2.2), Inches(0.2), name, size=11, bold=True, color=Theme.white)
        add_text(slide, card_x + Inches(0.12), Inches(1.98), Inches(2.2), Inches(0.22), period, size=10, bold=True, color=Theme.sub)
        add_bullets(slide, card_x + Inches(0.14), Inches(2.3), Inches(2.15), Inches(2.2), bullets, size=10)
        add_text(slide, card_x + Inches(0.14), Inches(4.7), Inches(2.15), Inches(0.2), "達成基準", size=10, bold=True, color=Theme.navy)
        add_text(
            slide,
            card_x + Inches(0.14),
            Inches(4.95),
            Inches(2.15),
            Inches(1.2),
            ["安定と適応", "認知療法修了", "関係修復", "就労内定", "経済的自立"][i],
            size=10,
            color=Theme.sub,
        )
    add_text(slide, Inches(10.8), Inches(1.12), Inches(2.2), Inches(0.2), "標準期間: 18ヶ月〜", size=10, bold=True, color=Theme.accent2)


def group_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "相模原ダルクグループ", "通所2拠点・入寮6拠点の包括ネットワーク", "11")
    add_rect(slide, Inches(0.55), Inches(1.38), Inches(12.2), Inches(0.8), Theme.card, Theme.panel)
    add_text(slide, Inches(0.8), Inches(1.58), Inches(8.5), Inches(0.3), "全8施設が連携し、状態に応じた切れ目ない回復支援を提供", size=13, bold=True, color=Theme.navy)

    add_rect(slide, Inches(0.62), Inches(2.35), Inches(6.0), Inches(1.8), Theme.card, Theme.panel)
    add_text(slide, Inches(0.84), Inches(2.52), Inches(5.5), Inches(0.25), "通所（2拠点）", size=12, bold=True, color=Theme.accent2)
    add_bullets(
        slide,
        Inches(0.84),
        Inches(2.82),
        Inches(5.6),
        Inches(1.18),
        [
            "相模原ダルク デイケアセンター: 各依存別プログラム / 相談受付・家族会",
            "相模原ダルク OTC: 就労継続支援B型 / 弁当惣菜・野菜づくり / 一般就労準備",
        ],
        size=11,
    )

    homes = [
        "大和PCC（初期）: 定員25名・断薬断酒特化",
        "町田RC（中期）: 欲求・身体症状が落ち着いた方",
        "愛川TC（中期）: バリアフリー完備",
        "上溝HRC（中期）: 心の癒しに特化",
        "相模原WPH（自立）: 個室9室・就労準備",
        "西門ACC（卒業後）: 仲間と暮らし自立を目指す",
    ]
    add_rect(slide, Inches(6.74), Inches(2.35), Inches(6.0), Inches(4.52), Theme.card, Theme.panel)
    add_text(slide, Inches(6.96), Inches(2.52), Inches(5.5), Inches(0.25), "入寮（6拠点）", size=12, bold=True, color=Theme.accent)
    y = Inches(2.84)
    for i, item in enumerate(homes):
        add_rect(slide, Inches(6.96 + (i % 2) * 2.88), y + Inches((i // 2) * 1.22), Inches(2.72), Inches(1.0), RGBColor(247, 250, 255), Theme.panel)
        add_text(
            slide,
            Inches(7.06 + (i % 2) * 2.88),
            y + Inches((i // 2) * 1.22) + Inches(0.1),
            Inches(2.52),
            Inches(0.82),
            item,
            size=9,
            color=Theme.text,
        )


def program_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "回復プログラム（デイ・ナイト・個別）", "集団・生活・個別支援を統合する3プログラム", "12")
    columns = [
        ("デイケア（通所）", Theme.accent2, ["ミーティング（言いっぱなし・聞きっぱなし）", "12ステップ", "SAGARPP（再発予防/CBT）", "ワークショップ・アート", "スポーツ＆プレジャー", "平日 9:00-17:00"]),
        ("ナイトケア（寮）", Theme.accent3, ["睡眠・食事改善", "集団生活での関係構築", "料理・洗濯の練習", "金銭・服薬管理", "24時間体制", "生活基盤の再構築"]),
        ("個別サポート", Theme.accent, ["個別カウンセリング", "セルフケア指導", "就労トレーニング", "継続的状態管理", "随時対応", "状態に合わせた継続支援"]),
    ]
    for i, (title, col, bullets) in enumerate(columns):
        x = Inches(0.72 + i * 4.2)
        add_rect(slide, x, Inches(1.55), Inches(3.95), Inches(5.22), Theme.card, Theme.panel)
        add_rect(slide, x, Inches(1.55), Inches(3.95), Inches(0.46), col)
        add_text(slide, x + Inches(0.18), Inches(1.68), Inches(3.5), Inches(0.2), title, size=13, bold=True, color=Theme.white)
        add_bullets(slide, x + Inches(0.18), Inches(2.12), Inches(3.6), Inches(4.45), bullets, size=11)


def features7_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "依存症の理解（7つの特徴）", "科学的理解に基づく回復支援の前提条件", "13")
    features = [
        "一次性の病気（原因は薬物使用そのもの）",
        "慢性の病気（生涯ケアが必要）",
        "進行性の病気（使用継続で喪失拡大）",
        "死亡率が高い（事故・自殺・ODリスク）",
        "性格が変化する（回復で本来性を回復）",
        "依存対象が移行（交差依存）",
        "人を巻き込む病気（家族支援が鍵）",
    ]
    add_rect(slide, Inches(0.62), Inches(1.42), Inches(12.1), Inches(0.62), Theme.card, Theme.panel)
    add_text(slide, Inches(0.85), Inches(1.6), Inches(11.5), Inches(0.22), "「意志の弱さ」ではなく「治療が必要な病気」として理解する", size=12, bold=True, color=Theme.navy)
    for i, item in enumerate(features):
        x = Inches(0.72 + (i % 2) * 6.2)
        y = Inches(2.24 + (i // 2) * 1.2)
        add_rect(slide, x, y, Inches(6.02), Inches(1.0), Theme.card, Theme.panel)
        add_text(slide, x + Inches(0.18), y + Inches(0.18), Inches(5.65), Inches(0.6), f"{i+1}. {item}", size=11, color=Theme.text)


def cross_addiction_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "交差依存（クロスアディクション）", "本命停止後に別依存へ移行しスリップへ戻る悪循環", "14")
    add_rect(slide, Inches(0.72), Inches(1.46), Inches(7.9), Inches(5.3), Theme.card, Theme.panel)
    add_text(slide, Inches(0.98), Inches(1.7), Inches(7.3), Inches(0.25), "悪循環の5ステップ", size=13, bold=True, color=Theme.navy)
    steps = [
        "1) 本命の依存をやめる",
        "2) 苦しみが増える（不安・渇望）",
        "3) 他の依存へ移行（処方薬・ギャンブル等）",
        "4) 問題解決できない",
        "5) 本命へ戻る（スリップ）",
    ]
    add_bullets(slide, Inches(0.98), Inches(2.05), Inches(7.4), Inches(2.6), steps, size=11)
    add_text(slide, Inches(0.98), Inches(4.95), Inches(7.4), Inches(0.22), "予防策: 早期介入 / 包括的再発予防 / ピア支援・12ステップ", size=11, bold=True, color=Theme.accent2)

    add_rect(slide, Inches(8.82), Inches(1.46), Inches(3.8), Inches(5.3), Theme.card, Theme.panel)
    add_text(slide, Inches(9.05), Inches(1.7), Inches(3.4), Inches(0.25), "依存タイプ", size=13, bold=True, color=Theme.navy)
    add_bullets(
        slide,
        Inches(9.05),
        Inches(2.05),
        Inches(3.3),
        Inches(3.6),
        ["物質依存: アルコール / 薬物", "プロセス依存: ギャンブル / 買い物 / ネット", "関係依存: 共依存 / 性依存", "対応: 包括的再発予防 + 家族支援"],
        size=10,
    )


def paws_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "長期離脱症状（PAWS）の理解", "12-24ヶ月寛解期における6症状と対処ポイント", "15")
    add_rect(slide, Inches(0.62), Inches(1.42), Inches(12.1), Inches(0.8), Theme.card, Theme.panel)
    add_text(slide, Inches(0.88), Inches(1.62), Inches(11.5), Inches(0.24), "ピーク 3-6ヶ月 / 寛解 12-24ヶ月 / 波を繰り返しながら徐々に回復", size=12, bold=True, color=Theme.navy)
    symptoms = [
        "① 睡眠障害（不眠・過眠・早朝覚醒）",
        "② ストレス過敏（音・視線・人混み）",
        "③ 記憶障害（短期記憶の低下）",
        "④ 感情障害（怒り・無感情・不安定）",
        "⑤ 身体バランス不調（疲労・めまい・頭痛）",
        "⑥ 思考障害（焦り・固執・集中力低下）",
    ]
    for i, item in enumerate(symptoms):
        x = Inches(0.72 + (i % 2) * 6.2)
        y = Inches(2.38 + (i // 2) * 1.35)
        add_rect(slide, x, y, Inches(6.02), Inches(1.15), Theme.card, Theme.panel)
        add_text(slide, x + Inches(0.2), y + Inches(0.2), Inches(5.6), Inches(0.7), item, size=11, color=Theme.text)
    add_text(slide, Inches(0.88), Inches(6.74), Inches(11.5), Inches(0.22), "重要: スリップリスクが高まる時期は一人で抱え込まずチームで対応", size=11, bold=True, color=RGBColor(201, 60, 50))


def craft_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "家族支援プログラム（CRAFT / 家族会）", "家族は回復の重要パートナー", "16")
    add_rect(slide, Inches(0.62), Inches(1.42), Inches(8.2), Inches(5.4), Theme.card, Theme.panel)
    add_text(slide, Inches(0.88), Inches(1.63), Inches(7.7), Inches(0.24), "CRAFT 5ステップ", size=13, bold=True, color=Theme.navy)
    steps = [
        "1. 理解・学ぶ（依存症理解とCRAFT基礎）",
        "2. 相談・家族会（毎週相談 / 月1家族会）",
        "3. 関わりを整える（肯定的対話 / 境界線設定）",
        "4. パターンを断つ（共依存・イネイブリング脱却）",
        "5. 継続と強化（回復行動を支える仕組み）",
    ]
    add_bullets(slide, Inches(0.88), Inches(1.95), Inches(7.7), Inches(3.9), steps, size=11)

    add_rect(slide, Inches(8.95), Inches(1.42), Inches(3.78), Inches(5.4), Theme.card, Theme.panel)
    add_text(slide, Inches(9.2), Inches(1.63), Inches(3.25), Inches(0.24), "援助の質を高める5点", size=13, bold=True, color=Theme.navy)
    add_bullets(
        slide,
        Inches(9.2),
        Inches(1.95),
        Inches(3.3),
        Inches(4.4),
        [
            "原因探しをしない",
            "責めない・なじらない",
            "あなたは一人じゃない",
            "タイミングがすべて",
            "イネイブリングを避ける",
        ],
        size=10,
    )


def network_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "連携機関とネットワーク", "医療・行政・司法・地域と接続し予防から社会復帰まで実装", "17")
    cards = [
        ("医療（MEDICAL）", ["北里大学病院（KIPP）", "高尾駒木野病院", "相模湖病院"]),
        ("行政（PUBLIC）", ["相模原市精神保健福祉センター（FLOW）", "東京多摩総合精神保健福祉センター（TAMARPP）"]),
        ("司法（JUSTICE）", ["横浜保護観察所", "府中刑務所", "八街少年院"]),
        ("地域（COMMUNITY）", ["地域行政（福祉課・保健所）", "専門病院", "学校", "企業（就労連携）"]),
    ]
    for i, (title, bullets) in enumerate(cards):
        x = Inches(0.72 + (i % 2) * 6.2)
        y = Inches(1.52 + (i // 2) * 2.67)
        add_rect(slide, x, y, Inches(6.02), Inches(2.45), Theme.card, Theme.panel)
        add_rect(slide, x, y, Inches(6.02), Inches(0.36), [Theme.accent2, Theme.accent, Theme.accent3, RGBColor(123, 92, 233)][i])
        add_text(slide, x + Inches(0.18), y + Inches(0.08), Inches(5.5), Inches(0.2), title, size=12, bold=True, color=Theme.white)
        add_bullets(slide, x + Inches(0.2), y + Inches(0.48), Inches(5.6), Inches(1.78), bullets, size=11)


def contact_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header(slide, "お問い合わせ｜CHANGE YOUR LIFE!", "依存症からの回復を全力で支えます", "18")
    add_rect(slide, Inches(0.58), Inches(1.45), Inches(12.2), Inches(1.3), Theme.card, Theme.panel)
    add_text(slide, Inches(0.95), Inches(1.78), Inches(11.5), Inches(0.42), "CHANGE YOUR LIFE!", size=36, bold=True, color=Theme.navy, align=PP_ALIGN.CENTER)
    add_text(slide, Inches(0.95), Inches(2.25), Inches(11.5), Inches(0.25), "変われる。変われた仲間がいる。", size=14, color=Theme.sub, align=PP_ALIGN.CENTER)

    cards = [
        ("電話でのご相談", ["042-707-0391", "平日 9:00-17:00 / 土祝〜14:00"]),
        ("お問い合わせフォーム", ["https://s-darc.com/", "24時間受付可能"]),
        ("所在地 / 運営", ["〒252-0237 相模原市中央区千代田3-3-20", "一般社団法人 相模原ダルク / 代表理事 田中秀泰"]),
    ]
    for i, (title, lines) in enumerate(cards):
        x = Inches(0.78 + i * 4.15)
        add_rect(slide, x, Inches(3.1), Inches(3.92), Inches(3.0), Theme.card, Theme.panel)
        add_rect(slide, x, Inches(3.1), Inches(3.92), Inches(0.42), [Theme.accent2, Theme.accent, Theme.accent3][i])
        add_text(slide, x + Inches(0.16), Inches(3.2), Inches(3.5), Inches(0.2), title, size=12, bold=True, color=Theme.white)
        add_bullets(slide, x + Inches(0.18), Inches(3.62), Inches(3.5), Inches(2.3), lines, size=11)


def build_part2(output_part2: Path):
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    cover(prs)
    stage_slide(prs)
    group_slide(prs)
    program_slide(prs)
    features7_slide(prs)
    cross_addiction_slide(prs)
    paws_slide(prs)
    craft_slide(prs)
    network_slide(prs)
    contact_slide(prs)

    output_part2.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_part2))


def verify_part2(output_part2: Path):
    prs = Presentation(str(output_part2))
    if len(prs.slides) != 10:
        raise RuntimeError(f"Unexpected part2 slide count: {len(prs.slides)} (expected 10)")
    required = [
        "5ステージ制 回復プログラム",
        "相模原ダルクグループ",
        "回復プログラム（デイ・ナイト・個別）",
        "依存症の理解（7つの特徴）",
        "交差依存",
        "PAWS",
        "CRAFT",
        "連携機関とネットワーク",
        "お問い合わせ",
    ]
    text = []
    for s in prs.slides:
        for sh in s.shapes:
            t = getattr(sh, "text", "")
            if t:
                text.append(t)
    joined = "\n".join(text)
    miss = [k for k in required if k not in joined]
    if miss:
        raise RuntimeError(f"Part2 missing required sections: {miss}")


def copy_slide_xml(dst: Presentation, src_slide):
    slide = dst.slides.add_slide(dst.slide_layouts[6])
    for shape in src_slide.shapes:
        new_el = copy.deepcopy(shape.element)
        slide.shapes._spTree.insert_element_before(new_el, "p:extLst")


def build_complete(part1_input: Path, part2_input: Path, output_complete: Path):
    if not part1_input.exists():
        raise FileNotFoundError(f"Part1 input not found: {part1_input}")
    p1 = Presentation(str(part1_input))
    p2 = Presentation(str(part2_input))

    merged = Presentation()
    merged.slide_width = p1.slide_width
    merged.slide_height = p1.slide_height

    for slide in p1.slides:
        copy_slide_xml(merged, slide)
    # Part2は表紙を除いて結合（全体で18枚に揃える）
    for slide in list(p2.slides)[1:]:
        copy_slide_xml(merged, slide)

    output_complete.parent.mkdir(parents=True, exist_ok=True)
    merged.save(str(output_complete))


def verify_complete(output_complete: Path):
    prs = Presentation(str(output_complete))
    if len(prs.slides) != 18:
        raise RuntimeError(
            f"Unexpected complete slide count: {len(prs.slides)} (expected 18)"
        )
    needed = [
        "本日のご説明内容",
        "日本の依存症情勢",
        "3段階予防",
        "組織概要",
        "実績ハイライト",
        "強み（3/3）",
        "5ステージ制 回復プログラム",
        "相模原ダルクグループ",
        "CRAFT",
        "連携機関とネットワーク",
        "CHANGE YOUR LIFE!",
    ]
    text = []
    for s in prs.slides:
        for sh in s.shapes:
            t = getattr(sh, "text", "")
            if t:
                text.append(t)
    joined = "\n".join(text)
    miss = [k for k in needed if k not in joined]
    if miss:
        raise RuntimeError(f"Complete deck missing sections: {miss}")


def main():
    args = parse_args()
    part1_input = Path(args.part1_input).expanduser().resolve()
    output_part2 = Path(args.output_part2).expanduser().resolve()
    output_complete = Path(args.output_complete).expanduser().resolve()

    build_part2(output_part2)
    verify_part2(output_part2)
    build_complete(part1_input, output_part2, output_complete)
    verify_complete(output_complete)
    print(f"Generated part2: {output_part2}")
    print(f"Generated complete: {output_complete}")


if __name__ == "__main__":
    main()
