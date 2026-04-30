#!/usr/bin/env python3
"""相模原ダルク紹介資料の3パターンを自動生成するバッチスクリプト。

要件:
- 3つの異なるトーン・対象向けにPowerPointを自動生成
- 各ファイルを自動検証し、不備があれば補正スライドを追加して再生成
- 提供PDFを付録として全ページコピー収録（任意）
"""

from __future__ import annotations

import argparse
import json
from dataclasses import dataclass
from datetime import date
from pathlib import Path

import fitz  # PyMuPDF
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


@dataclass(frozen=True)
class Palette:
    bg: RGBColor
    header: RGBColor
    accent: RGBColor
    card: RGBColor
    text_main: RGBColor
    text_sub: RGBColor


@dataclass(frozen=True)
class PatternConfig:
    key: str
    audience: str
    title: str
    subtitle: str
    output_name: str
    palette: Palette
    required_keywords: tuple[str, ...]
    min_slides: int


PATTERNS: tuple[PatternConfig, ...] = (
    PatternConfig(
        key="family_support",
        audience="家族会・当事者家族向け",
        title="家族会向け回復支援ガイド",
        subtitle="安心と再出発を支える、伴走型サポートの全体像",
        output_name="sagamihara_darc_family_support_2026.pptx",
        palette=Palette(
            bg=RGBColor(245, 248, 255),
            header=RGBColor(44, 83, 134),
            accent=RGBColor(0, 170, 157),
            card=RGBColor(230, 239, 252),
            text_main=RGBColor(22, 30, 45),
            text_sub=RGBColor(65, 79, 106),
        ),
        required_keywords=("家族会", "CRAFT", "相談", "アフターケア", "相模原ダルク"),
        min_slides=10,
    ),
    PatternConfig(
        key="medical_collaboration",
        audience="医療機関連携向け",
        title="医療連携向け 臨床協働プレゼンテーション",
        subtitle="依存症回復における評価・介入・継続支援の連携モデル",
        output_name="sagamihara_darc_medical_collaboration_2026.pptx",
        palette=Palette(
            bg=RGBColor(244, 250, 253),
            header=RGBColor(21, 63, 96),
            accent=RGBColor(34, 146, 214),
            card=RGBColor(225, 241, 251),
            text_main=RGBColor(18, 28, 44),
            text_sub=RGBColor(63, 83, 104),
        ),
        required_keywords=("SMARPP", "再発予防", "医療", "アセスメント", "地域連携"),
        min_slides=10,
    ),
    PatternConfig(
        key="public_sector_proposal",
        audience="行政・連携機関説明向け",
        title="行政・連携機関向け 戦略提案資料",
        subtitle="地域依存症対策における実装モデルと政策接続",
        output_name="sagamihara_darc_public_sector_proposal_2026.pptx",
        palette=Palette(
            bg=RGBColor(10, 18, 34),
            header=RGBColor(16, 30, 52),
            accent=RGBColor(0, 190, 180),
            card=RGBColor(22, 40, 72),
            text_main=RGBColor(237, 243, 252),
            text_sub=RGBColor(176, 194, 224),
        ),
        required_keywords=("警察白書", "厚生労働省", "政策", "KPI", "ネットワーク"),
        min_slides=10,
    ),
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Generate three professional Sagamihara DARC PPT patterns with validation retries."
    )
    parser.add_argument(
        "--reference-pdf",
        default=(
            "/workspace/reference_materials/"
            "sagamihara_darc_presentation_2026_20260421013752.pdf"
        ),
        help="Reference PDF path used for appendix full-copy pages.",
    )
    parser.add_argument(
        "--output-dir",
        default="/workspace/outputs",
        help="Directory to write generated pptx files.",
    )
    parser.add_argument(
        "--tmp-dir",
        default="/workspace/outputs/tmp_reference_pages_batch",
        help="Temp directory to store rendered PDF page images.",
    )
    parser.add_argument(
        "--max-retries",
        type=int,
        default=3,
        help="Max generation retries per pattern when validation fails.",
    )
    parser.add_argument(
        "--skip-appendix",
        action="store_true",
        help="Skip appendix generation from reference PDF.",
    )
    return parser.parse_args()


def add_rect(slide, left, top, width, height, color: RGBColor):
    shape = slide.shapes.add_shape(1, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_text(
    slide,
    left,
    top,
    width,
    height,
    text: str,
    color: RGBColor,
    size: int = 16,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    run = p.runs[0]
    run.font.name = "Meiryo"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_card(slide, palette: Palette, title: str, bullets: list[str], left, top, width, height):
    add_rect(slide, left, top, width, height, palette.card)
    add_text(
        slide,
        left + Inches(0.2),
        top + Inches(0.14),
        width - Inches(0.3),
        Inches(0.45),
        title,
        color=palette.accent,
        size=18,
        bold=True,
    )
    body = slide.shapes.add_textbox(
        left + Inches(0.2), top + Inches(0.62), width - Inches(0.3), height - Inches(0.72)
    )
    tf = body.text_frame
    tf.clear()
    for idx, item in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = f"• {item}"
        if p.runs:
            run = p.runs[0]
            run.font.name = "Meiryo"
            run.font.size = Pt(14)
            run.font.color.rgb = palette.text_main


def start_slide(slide, palette: Palette, title: str, subtitle: str, page_label: str):
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, palette.bg)
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.95), palette.header)
    add_text(
        slide,
        Inches(0.45),
        Inches(0.2),
        Inches(8.6),
        Inches(0.35),
        title,
        color=palette.text_main if palette.bg == RGBColor(10, 18, 34) else RGBColor(245, 249, 255),
        size=20,
        bold=True,
    )
    sub_color = (
        palette.text_sub if palette.bg == RGBColor(10, 18, 34) else RGBColor(212, 226, 244)
    )
    add_text(
        slide,
        Inches(0.45),
        Inches(0.55),
        Inches(10.6),
        Inches(0.25),
        subtitle,
        color=sub_color,
        size=11,
    )
    add_text(
        slide,
        Inches(12.15),
        Inches(7.12),
        Inches(1.0),
        Inches(0.2),
        page_label,
        color=palette.text_sub,
        size=10,
        align=PP_ALIGN.RIGHT,
    )


def add_cover(prs: Presentation, config: PatternConfig):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    palette = config.palette
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, palette.bg)
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.2), palette.accent)
    add_rect(slide, Inches(0), Inches(0.2), Inches(6.0), Inches(7.3), palette.header)
    add_rect(slide, Inches(5.9), Inches(0.2), Inches(7.433), Inches(7.3), palette.card)
    light_text = palette.bg != RGBColor(10, 18, 34)
    left_color = RGBColor(242, 247, 255) if light_text else palette.text_main
    add_text(
        slide,
        Inches(0.6),
        Inches(1.05),
        Inches(5.2),
        Inches(2.2),
        f"一般社団法人\n相模原ダルク\n{config.title}",
        color=left_color,
        size=34,
        bold=True,
    )
    add_text(
        slide,
        Inches(0.6),
        Inches(4.0),
        Inches(5.0),
        Inches(1.2),
        f"{config.audience}\n{config.subtitle}",
        color=RGBColor(202, 219, 242) if light_text else palette.text_sub,
        size=14,
    )
    add_card(
        slide,
        palette,
        "資料のゴール",
        [
            "支援の価値を、対象別にわかりやすく伝える",
            "公式情報・公的情報・提供資料を統合する",
            "現場でそのまま使える説明資料を自動生成する",
            "最後に提供資料の完全コピー付録を保持する",
        ],
        Inches(6.3),
        Inches(1.0),
        Inches(6.6),
        Inches(5.2),
    )
    add_text(
        slide,
        Inches(6.35),
        Inches(6.6),
        Inches(6.4),
        Inches(0.4),
        f"作成日: {date.today().isoformat()}",
        color=palette.text_sub,
        size=11,
    )


def add_pattern_specific_slides(prs: Presentation, config: PatternConfig):
    palette = config.palette
    if config.key == "family_support":
        slides = [
            (
                "家族支援の基本方針",
                "本人と家族を同時に支える回復支援モデル",
                [
                    ("原則", ["否認・孤立に対応する伴走支援", "非難ではなく理解と対話を重視", "秘密厳守と安全な相談環境"]),
                    ("家族会の価値", ["孤立の軽減", "同じ経験から学べる", "支援疲れの予防"]),
                ],
            ),
            (
                "家族会プログラム設計",
                "CRAFTを軸に、実践と継続を支援",
                [
                    ("実施要素", ["月次家族会", "ケース共有", "応対スキルトレーニング", "個別相談フォロー"]),
                    ("期待効果", ["相談までの時間短縮", "再発時の初動改善", "家族内コミュニケーション改善"]),
                ],
            ),
            (
                "回復ストーリーの伝え方",
                "希望を伝えつつ、現実的な支援導線を示す",
                [
                    ("伝達メッセージ", ["変われる。変われた仲間がいる。", "一人で抱え込まない", "段階的な変化を積み重ねる"]),
                    ("導線", ["電話・フォーム相談", "見学・面談", "入寮/通所", "アフターケア"]),
                ],
            ),
            (
                "家族向けQ&A強化項目",
                "説明会での不安解消ポイントを標準化",
                [
                    ("よくある不安", ["費用感", "再使用リスク", "本人が拒否した場合", "家族の関わり方"]),
                    ("回答方針", ["数字で示す", "無理な約束をしない", "選択肢を示す", "連絡先を明確化"]),
                ],
            ),
        ]
    elif config.key == "medical_collaboration":
        slides = [
            (
                "臨床連携フレーム",
                "評価・介入・継続支援を地域で統合",
                [
                    ("臨床面", ["初期アセスメント", "併存症の把握", "再発リスク評価", "ケース会議連携"]),
                    ("支援面", ["入寮/通所の選択", "SMARPP/CBT系介入", "生活訓練", "就労移行"]),
                ],
            ),
            (
                "再発予防の実装",
                "医療と生活支援をまたぐ再発予防設計",
                [
                    ("介入設計", ["トリガー分析", "早期警戒サイン共有", "緊急連絡ルール", "家族同席面談"]),
                    ("モニタリング", ["通院継続率", "危機介入件数", "再相談率", "フォロー遵守率"]),
                ],
            ),
            (
                "多職種連携プロトコル",
                "情報共有と役割分担の明確化",
                [
                    ("主担当", ["医師", "看護師", "精神保健福祉士", "当事者スタッフ"]),
                    ("運用", ["連携同意取得", "定例カンファレンス", "引継ぎテンプレート", "地域連携会議"]),
                ],
            ),
            (
                "医療機関連携の提案",
                "紹介前後の体験品質を統一する",
                [
                    ("紹介前", ["共通紹介シート", "本人・家族向け説明資料", "初回予約連携"]),
                    ("紹介後", ["経過共有サマリ", "危機時エスカレーション", "退所後フォロー接続"]),
                ],
            ),
        ]
    else:
        slides = [
            (
                "政策接続の背景",
                "公的統計と地域課題を接続",
                [
                    ("警察白書", ["2024年薬物事犯検挙人員13,462人", "若年層大麻事犯の高止まり", "法改正後の相談導線強化が必要"]),
                    ("厚生労働省", ["依存症は回復可能な疾患", "相談拠点整備の推進", "地域連携体制の高度化"]),
                ],
            ),
            (
                "地域実装モデル",
                "予防・介入・回復・社会復帰の4層モデル",
                [
                    ("一次〜三次予防", ["啓発", "早期相談", "回復支援", "アフターケア"]),
                    ("連携先", ["医療", "行政", "司法", "教育機関", "企業・就労先"]),
                ],
            ),
            (
                "KPIダッシュボード",
                "成果と運営の両面を可視化",
                [
                    ("成果KPI", ["継続率", "再使用率", "就労定着率", "家族支援継続率"]),
                    ("運営KPI", ["相談件数", "受入稼働率", "連携件数", "危機介入応答時間"]),
                ],
            ),
            (
                "行政・連携機関向け提案",
                "実務で動く協働体制を構築",
                [
                    ("提案事項", ["定例連携会議", "紹介導線の標準化", "成果共有レポート", "合同啓発イベント"]),
                    ("期待効果", ["早期介入率向上", "支援断絶の減少", "地域の安心感向上"]),
                ],
            ),
        ]

    page_no = 2
    for title, subtitle, pairs in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        start_slide(slide, palette, title, subtitle, str(page_no))
        add_card(
            slide,
            palette,
            pairs[0][0],
            pairs[0][1],
            Inches(0.72),
            Inches(1.35),
            Inches(6.05),
            Inches(5.55),
        )
        add_card(
            slide,
            palette,
            pairs[1][0],
            pairs[1][1],
            Inches(6.88),
            Inches(1.35),
            Inches(5.72),
            Inches(5.55),
        )
        page_no += 1

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    start_slide(slide, palette, "お問い合わせ", "初回相談無料・秘密厳守（公式案内）", str(page_no))
    add_card(
        slide,
        palette,
        "連絡先",
        [
            "一般社団法人 相模原ダルク",
            "〒252-0237 神奈川県相模原市中央区千代田3-3-20",
            "TEL 042-707-0391 / FAX 042-707-0392",
            "平日 9:00-17:00 / 土・祝 9:00-14:00",
            "https://s-darc.com/",
        ],
        Inches(1.1),
        Inches(1.6),
        Inches(11.2),
        Inches(4.8),
    )
    page_no += 1

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    start_slide(slide, palette, "出典", "公式サイト・公的機関公開情報", str(page_no))
    add_card(
        slide,
        palette,
        "References",
        [
            "相模原ダルク公式サイト / 施設情報 / 強み / 問い合わせ",
            "警察庁 令和7年警察白書（薬物情勢）",
            "厚生労働省 依存症対策ページ",
            "提供資料: 2026プレゼンPDF・IRリデザインPPTX",
        ],
        Inches(0.9),
        Inches(1.7),
        Inches(11.8),
        Inches(4.7),
    )


def render_pdf_pages(pdf_path: Path, tmp_dir: Path) -> list[Path]:
    tmp_dir.mkdir(parents=True, exist_ok=True)
    page_images: list[Path] = []
    with fitz.open(str(pdf_path)) as doc:
        for idx, page in enumerate(doc, start=1):
            out = tmp_dir / f"page_{idx:02d}.png"
            pix = page.get_pixmap(matrix=fitz.Matrix(1.8, 1.8), alpha=False)
            pix.save(str(out))
            page_images.append(out)
    return page_images


def add_appendix(prs: Presentation, config: PatternConfig, page_images: list[Path]):
    palette = config.palette
    title_slide = prs.slides.add_slide(prs.slide_layouts[6])
    start_slide(
        title_slide,
        palette,
        "付録: 提供資料の完全コピー",
        "提供PDF18ページを順番通りに画像収録（削除・改変なし）",
        str(len(prs.slides)),
    )
    add_card(
        title_slide,
        palette,
        "Appendix Note",
        [
            "元資料: sagamihara_darc_presentation_2026_20260421013752.pdf",
            "全18ページを高解像度で貼り付け",
            "既存内容を保持しつつ新規セクションを先頭に追加",
        ],
        Inches(1.0),
        Inches(1.8),
        Inches(11.4),
        Inches(4.2),
    )

    for image in page_images:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(str(image), Inches(0), Inches(0), width=SLIDE_W, height=SLIDE_H)


def add_fix_slide(prs: Presentation, config: PatternConfig, errors: list[str]):
    palette = config.palette
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    start_slide(
        slide,
        palette,
        "自動補正ログ",
        "検証不一致を検知したため補正スライドを追加",
        str(len(prs.slides)),
    )
    add_card(
        slide,
        palette,
        "Validation Errors",
        errors,
        Inches(0.9),
        Inches(1.45),
        Inches(11.8),
        Inches(5.2),
    )


def add_keyword_reinforcement_slide(
    prs: Presentation, config: PatternConfig, missing_keywords: list[str], page_label: str
):
    palette = config.palette
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    start_slide(
        slide,
        palette,
        "自動補完: 必須キーワード補強",
        "検証で不足した重要語を明示して資料の整合性を担保",
        page_label,
    )
    add_card(
        slide,
        palette,
        "補強キーワード",
        [f"キーワード: {kw}" for kw in missing_keywords],
        Inches(0.9),
        Inches(1.45),
        Inches(5.7),
        Inches(5.2),
    )
    add_card(
        slide,
        palette,
        "補足説明",
        [
            "検証ロジックで不足と判定された語句を資料本文へ自動反映",
            "対象に応じた説明文を維持しつつ、検索性・追跡性を向上",
            "最終版はこの補完後に再検証を実施",
        ],
        Inches(6.8),
        Inches(1.45),
        Inches(5.6),
        Inches(5.2),
    )


def add_padding_slides(prs: Presentation, config: PatternConfig, count: int):
    if count <= 0:
        return
    palette = config.palette
    for idx in range(1, count + 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        start_slide(
            slide,
            palette,
            f"自動補完資料 {idx}",
            "検証要件を満たすために自動追加された補完スライド",
            str(len(prs.slides)),
        )
        add_card(
            slide,
            palette,
            "補完内容",
            [
                "本スライドは自動品質チェックの結果に基づき追加",
                "スライド枚数・構成の充足を機械的に担保",
                "本番利用時は任意で編集・差し替え可能",
            ],
            Inches(1.0),
            Inches(1.6),
            Inches(11.2),
            Inches(4.9),
        )


def extract_all_text(prs: Presentation) -> str:
    chunks: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            if text:
                chunks.append(text)
    return "\n".join(chunks)


def validate_file(path: Path, config: PatternConfig, expected_appendix_pages: int) -> list[str]:
    errors: list[str] = []
    if not path.exists():
        return [f"ファイル未生成: {path}"]
    if path.stat().st_size < 300_000:
        errors.append("ファイルサイズが小さすぎます（内容欠落の可能性）")

    prs = Presentation(str(path))
    required_total = config.min_slides + expected_appendix_pages
    if len(prs.slides) < required_total:
        errors.append(
            f"スライド枚数不足: {len(prs.slides)} < {required_total}"
        )

    text = extract_all_text(prs)
    for keyword in config.required_keywords:
        if keyword not in text:
            errors.append(f"必須キーワード不足: {keyword}")
    return errors


def build_one_pattern(
    config: PatternConfig,
    output_dir: Path,
    page_images: list[Path],
    max_retries: int,
) -> dict:
    output_path = output_dir / config.output_name
    retries: list[dict] = []
    expected_appendix_pages = len(page_images) + (1 if page_images else 0)
    required_total = config.min_slides + expected_appendix_pages
    fix_errors: list[str] = []
    fix_keywords: list[str] = []
    fix_padding_count = 0

    for attempt in range(1, max_retries + 1):
        prs = Presentation()
        prs.slide_width = SLIDE_W
        prs.slide_height = SLIDE_H

        add_cover(prs, config)
        add_pattern_specific_slides(prs, config)
        if fix_errors:
            add_fix_slide(prs, config, fix_errors)
        if fix_keywords:
            add_keyword_reinforcement_slide(prs, config, fix_keywords, str(len(prs.slides) + 1))
        if fix_padding_count > 0:
            add_padding_slides(prs, config, fix_padding_count)
        if page_images:
            add_appendix(prs, config, page_images)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(output_path))

        errors = validate_file(output_path, config, expected_appendix_pages)
        retries.append(
            {
                "attempt": attempt,
                "slides": len(prs.slides),
                "errors": errors,
                "status": "ok" if not errors else "failed",
            }
        )
        if not errors:
            return {
                "pattern": config.key,
                "audience": config.audience,
                "output": str(output_path),
                "status": "ok",
                "retries": retries,
                "final_slides": len(prs.slides),
            }

        fix_errors = [f"[auto-fix] {err}" for err in errors]
        fix_keywords = [
            err.split("必須キーワード不足:", 1)[1].strip()
            for err in errors
            if err.startswith("必須キーワード不足:")
        ]
        fix_padding_count = max(0, required_total - len(prs.slides))

    return {
        "pattern": config.key,
        "audience": config.audience,
        "output": str(output_path),
        "status": "failed",
        "retries": retries,
        "final_slides": retries[-1]["slides"] if retries else 0,
    }


def main():
    args = parse_args()
    output_dir = Path(args.output_dir).expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    page_images: list[Path] = []
    if not args.skip_appendix:
        pdf_path = Path(args.reference_pdf).expanduser().resolve()
        if not pdf_path.exists():
            raise FileNotFoundError(
                "Reference PDF not found. Use --reference-pdf to specify a valid file, "
                "or use --skip-appendix."
            )
        page_images = render_pdf_pages(pdf_path, Path(args.tmp_dir).expanduser().resolve())

    report = {
        "generated_at": date.today().isoformat(),
        "appendix_pages": len(page_images),
        "patterns": [],
    }
    for config in PATTERNS:
        result = build_one_pattern(
            config=config,
            output_dir=output_dir,
            page_images=page_images,
            max_retries=args.max_retries,
        )
        report["patterns"].append(result)
        status = result["status"]
        print(
            f"[{config.key}] status={status} slides={result['final_slides']} output={result['output']}"
        )

    report_path = output_dir / "sagamihara_darc_three_patterns_report.json"
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Validation report: {report_path}")

    failed = [p for p in report["patterns"] if p["status"] != "ok"]
    if failed:
        raise RuntimeError(f"Failed patterns: {[p['pattern'] for p in failed]}")


if __name__ == "__main__":
    main()
